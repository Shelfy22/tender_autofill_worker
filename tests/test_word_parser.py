from pathlib import Path

from docx import Document

from app.config import Settings
from app.services.parsers.word import extract_word_text
from app.services.products import extract_deterministic_positions


def test_word_parser_extracts_nested_procurement_table(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("\u0422\u0435\u0445\u043d\u0438\u0447\u0435\u0441\u043a\u0430\u044f \u0434\u043e\u043a\u0443\u043c\u0435\u043d\u0442\u0430\u0446\u0438\u044f " * 30)
    outer = document.add_table(rows=1, cols=1)
    outer.cell(0, 0).text = "\u041e\u0431\u043e\u0441\u043d\u043e\u0432\u0430\u043d\u0438\u0435 \u041d\u041c\u0426\u0414"
    nested = outer.cell(0, 0).add_table(rows=2, cols=4)
    nested.rows[0].cells[0].text = "\u041e\u0431\u044a\u0435\u043a\u0442 \u0437\u0430\u043a\u0443\u043f\u043a\u0438"
    nested.rows[0].cells[1].text = "\u041a\u043e\u043b\u0438\u0447\u0435\u0441\u0442\u0432\u043e"
    nested.rows[0].cells[2].text = "\u0415\u0434. \u0438\u0437\u043c."
    nested.rows[0].cells[3].text = "\u0422\u0435\u0445\u043d\u0438\u0447\u0435\u0441\u043a\u0438\u0435 \u0445\u0430\u0440\u0430\u043a\u0442\u0435\u0440\u0438\u0441\u0442\u0438\u043a\u0438"
    nested.rows[1].cells[0].text = "\u0411\u043e\u043b\u0442 12x65"
    nested.rows[1].cells[1].text = "1"
    nested.rows[1].cells[2].text = "\u043a\u0433"
    nested.rows[1].cells[3].text = "\u0413\u041e\u0421\u0422 7798-70"
    path = tmp_path / "nested.docx"
    document.save(path)

    text, status, warnings = extract_word_text(
        path,
        "docx",
        Settings(postgres_dsn="postgresql://user:pass@localhost/db"),
    )
    positions = extract_deterministic_positions(text)

    assert status == "ok"
    assert warnings == []
    assert text.count("\u0422\u0430\u0431\u043b\u0438\u0446\u0430 Word 2") == 1
    assert [(item.product, item.quantity, item.unit) for item in positions] == [
        ("\u0411\u043e\u043b\u0442 12x65", 1.0, "\u043a\u0433")
    ]


def test_word_parser_preserves_section_and_table_order(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("4. ТЕХНИЧЕСКОЕ ЗАДАНИЕ")
    technical = document.add_table(rows=1, cols=1)
    technical.cell(0, 0).text = "Характеристики товара"
    document.add_paragraph("5. ОБОСНОВАНИЕ НАЧАЛЬНОЙ МАКСИМАЛЬНОЙ ЦЕНЫ")
    price = document.add_table(rows=1, cols=1)
    price.cell(0, 0).text = "Перечень товаров и цены"
    document.add_paragraph("Техническая документация " * 30)
    path = tmp_path / "composite.docx"
    document.save(path)

    text, status, _ = extract_word_text(
        path,
        "docx",
        Settings(postgres_dsn="postgresql://user:pass@localhost/db"),
    )

    assert status == "ok"
    assert text.index("4. ТЕХНИЧЕСКОЕ ЗАДАНИЕ") < text.index("Таблица Word 1")
    assert text.index("Таблица Word 1") < text.index(
        "5. ОБОСНОВАНИЕ НАЧАЛЬНОЙ МАКСИМАЛЬНОЙ ЦЕНЫ"
    )
    assert text.index("5. ОБОСНОВАНИЕ") < text.index("Таблица Word 2")
